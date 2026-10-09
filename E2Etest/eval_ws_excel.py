"""Batch-eval /realtime/stream against tts-50.xlsx + 1.wav … 50.wav.

Records hypothesis, CER, and client latency. One process, one wav at a time:
connect → PCM → stop → final → close → next. A second eval exits immediately.

    conda activate lingting
    pip install openpyxl          # once, Excel reader
    # server must be running: python run.py

    python E2Etest/eval_ws_excel.py
    python E2Etest/eval_ws_excel.py --dry-run
    python E2Etest/eval_ws_excel.py --from 1 --to 3
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import json
import os
import re
import sys
import time
from datetime import datetime
from math import gcd
from pathlib import Path

import httpx
import numpy as np
import soundfile as sf
from rapidfuzz.distance import Levenshtein
from scipy.signal import resample_poly
import websockets

ROOT = Path(__file__).resolve().parents[1]
_E2E = Path(__file__).resolve().parent
if str(_E2E) not in sys.path:
    sys.path.insert(0, str(_E2E))
from ws_envelope import FORMAT_16K, ws_audio_msg
AUDIO_DIR = Path(r"E:\cursorProject\tts_ws_python3_demo\audio")
XLSX_NAME = "tts-50.xlsx"
BASE_URL = "http://127.0.0.1:8765"
WS_URL = "ws://127.0.0.1:8765/realtime/stream"
TARGET_SR = 16000
LOCK_PATH = ROOT / "E2Etest" / "out" / "eval_ws.lock"

_HEADER_SEQ = {"序号", "id", "index", "seq", "no", "编号"}
_HEADER_REF = {"文本", "text", "ref", "reference", "内容"}


def resolve_server_log() -> Path:
    log_dir = ROOT / "logs"
    raw = os.getenv("LOG_DIR", "").strip()
    if raw:
        p = Path(raw).expanduser()
        log_dir = p if p.is_absolute() else (ROOT / p).resolve()
    today = log_dir / f"server_{datetime.now():%Y-%m-%d}.log"
    if today.is_file():
        return today
    legacy = log_dir / "server.log"
    if legacy.is_file():
        return legacy
    dated = sorted(log_dir.glob("server_*.log"))
    return dated[-1] if dated else today


def seq_stem(seq) -> str:
    if isinstance(seq, int):
        return str(seq)
    raw = str(seq).strip()
    try:
        f = float(raw)
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return raw


def find_wav(folder: Path, seq) -> Path | None:
    stem = seq_stem(seq)
    p = folder / f"{stem}.wav"
    return p if p.is_file() else None


def read_excel(path: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    start = 0
    seq_col, ref_col = 0, 1
    first = [str(c).strip() if c is not None else "" for c in rows[0]]
    if first and first[0] in _HEADER_SEQ:
        start = 1
        for i, name in enumerate(first):
            if name in _HEADER_SEQ:
                seq_col = i
            if name in _HEADER_REF:
                ref_col = i
    out: list[tuple[str, str]] = []
    for row in rows[start:]:
        if not row or row[seq_col] is None:
            continue
        ref = row[ref_col] if ref_col < len(row) else ""
        out.append((seq_stem(row[seq_col]), str(ref or "").strip()))
    return out


def normalize(text: str, strip_punct: bool) -> str:
    s = (text or "").strip().lower().replace(" ", "").replace("\u3000", "")
    if strip_punct:
        s = re.sub(r"[^\w\u4e00-\u9fff]+", "", s, flags=re.UNICODE)
    return s


def cer(ref: str, hyp: str, strip_punct: bool) -> float:
    r, h = normalize(ref, strip_punct), normalize(hyp, strip_punct)
    if not r:
        return 0.0 if not h else 1.0
    return Levenshtein.distance(r, h) / len(r)


def load_pcm_s16le(path: Path) -> tuple[bytes, float]:
    audio, sr = sf.read(str(path), always_2d=False)
    if getattr(audio, "ndim", 1) > 1:
        audio = audio.mean(axis=1)
    audio = np.asarray(audio, dtype=np.float32)
    if int(sr) != TARGET_SR:
        g = gcd(int(sr), TARGET_SR)
        audio = resample_poly(audio, TARGET_SR // g, int(sr) // g).astype(np.float32)
        sr = TARGET_SR
    peak = float(np.max(np.abs(audio))) if audio.size else 0.0
    if peak > 1.0:
        audio = audio / peak
    pcm = np.clip(audio * 32767.0, -32768, 32767).astype(np.int16).tobytes()
    return pcm, len(audio) / float(sr)


def parse_kv(line: str) -> dict[str, str]:
    if "WS infer_done |" not in line:
        return {}
    body = line.split("WS infer_done |", 1)[1]
    return {m.group(1): m.group(2) for m in re.finditer(r"(\w+)=(\S+)", body)}


def attach_log_metrics(rows: list[dict], log_path: Path) -> None:
    if not log_path.is_file():
        return
    by_sid: dict[str, dict[str, str]] = {}
    for raw in log_path.read_text(encoding="utf-8", errors="replace").splitlines():
        fields = parse_kv(raw)
        sid = fields.get("session_id")
        kind = fields.get("kind")
        if not sid or kind not in ("final", "speculative"):
            continue
        if fields.get("cancelled") == "1":
            continue
        by_sid[sid] = fields
    for row in rows:
        fields = by_sid.get(row.get("session_id") or "")
        if not fields:
            continue
        row["log_kind"] = fields.get("kind", "")
        row["log_audio_s"] = fields.get("audio", "").rstrip("s")
        row["log_queue_wait_ms"] = fields.get("queue_wait_ms", "")
        row["log_infer_ms"] = fields.get("infer_ms", "")
        row["log_encode_ms"] = fields.get("encode_ms", "")
        row["log_decode_ms"] = fields.get("decode_ms", "")
        row["log_n_chunks"] = fields.get("n_chunks", "")
        row["log_decode_skipped"] = fields.get("decode_skipped", "")
        row["log_decode_chunks"] = fields.get("decode_chunks", "")


def _pid_alive(pid: int) -> bool:
    if pid <= 0:
        return False
    if sys.platform == "win32":
        import ctypes

        handle = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)
        if handle:
            ctypes.windll.kernel32.CloseHandle(handle)
            return True
        return False
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def acquire_eval_lock() -> int:
    """Only one eval process may use the single server WebSocket."""
    LOCK_PATH.parent.mkdir(parents=True, exist_ok=True)
    for _ in range(2):
        try:
            fd = os.open(str(LOCK_PATH), os.O_CREAT | os.O_EXCL | os.O_RDWR)
            os.write(fd, str(os.getpid()).encode("ascii"))
            return fd
        except FileExistsError:
            raw = LOCK_PATH.read_text(encoding="utf-8", errors="replace").strip()
            old = int(raw) if raw.isdigit() else 0
            if old and _pid_alive(old):
                print(
                    f"another eval is already running (pid {old}). "
                    "Run one eval at a time.",
                    file=sys.stderr,
                )
                raise SystemExit(2)
            try:
                LOCK_PATH.unlink()
            except OSError:
                pass
    print("could not acquire eval lock", file=sys.stderr)
    raise SystemExit(2)


def release_eval_lock(fd: int) -> None:
    try:
        os.close(fd)
    except OSError:
        pass
    try:
        LOCK_PATH.unlink()
    except OSError:
        pass


async def connect_session(session_id: str, busy_wait_s: float):
    """Retry while the server still holds the previous WebSocket (WS_BUSY)."""
    url = f"{WS_URL}?request_id={session_id}"
    deadline = time.perf_counter() + max(0.5, busy_wait_s)
    last = None
    while time.perf_counter() < deadline:
        ws = await websockets.connect(url, max_size=None, open_timeout=10)
        raw = await ws.recv()
        hello = json.loads(raw.decode("utf-8") if isinstance(raw, bytes) else raw)
        if hello.get("status") == "connected":
            return ws, hello
        last = hello
        try:
            await ws.close()
        except Exception:
            pass
        if hello.get("code") != "WS_BUSY":
            break
        await asyncio.sleep(0.4)
    raise RuntimeError(f"handshake failed: {last}")


async def transcribe_one(
    pcm: bytes,
    *,
    session_id: str,
    realtime: bool,
    chunk_ms: int,
    language: str,
    timeout_s: float,
    busy_wait_s: float = 20.0,
) -> dict:
    chunk = max(2, int(TARGET_SR * chunk_ms / 1000) * 2)
    t0 = time.perf_counter()
    wall_start = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    first_partial = None
    t_last_send = None
    t_final = None
    final_text = ""
    eos_reason = ""
    n_partial = 0
    hello_sid = session_id

    ws, hello = await connect_session(session_id, busy_wait_s)
    hello_sid = str(hello.get("session_id") or session_id)
    try:
        closed = asyncio.Event()

        async def reader():
            nonlocal first_partial, t_final, final_text, eos_reason, n_partial
            try:
                while True:
                    raw = await ws.recv()
                    if isinstance(raw, bytes):
                        raw = raw.decode("utf-8")
                    msg = json.loads(raw)
                    kind = msg.get("type")
                    if kind == "partial":
                        n_partial += 1
                        if first_partial is None:
                            first_partial = time.perf_counter()
                    if kind == "final" or msg.get("is_final"):
                        final_text = msg.get("text") or ""
                        eos_reason = str(msg.get("eos_reason") or "")
                        t_final = time.perf_counter()
                    if msg.get("status") == "closed":
                        if not eos_reason:
                            eos_reason = str(msg.get("reason") or "")
                        closed.set()
                        return
            except websockets.ConnectionClosed:
                closed.set()

        task = asyncio.create_task(reader())
        try:
            offset = 0
            first = True
            while offset < len(pcm) and not closed.is_set():
                piece = pcm[offset : offset + chunk]
                try:
                    await ws.send(ws_audio_msg(0 if first else 1, piece, FORMAT_16K))
                except websockets.ConnectionClosed:
                    break
                first = False
                t_last_send = time.perf_counter()
                offset += chunk
                if realtime:
                    await asyncio.sleep(chunk_ms / 1000.0)
            if not closed.is_set():
                try:
                    await ws.send(ws_audio_msg(2))
                except websockets.ConnectionClosed:
                    pass
            await asyncio.wait_for(closed.wait(), timeout=timeout_s)
        finally:
            task.cancel()
            try:
                await task
            except asyncio.CancelledError:
                pass
    finally:
        try:
            await ws.close()
        except Exception:
            pass

    t_end = time.perf_counter()
    last = t_last_send or t0
    final_at = t_final or t_end
    return {
        "session_id": hello_sid,
        "t_start": wall_start,
        "hyp": final_text,
        "eos_reason": eos_reason,
        "n_partial": n_partial,
        "first_partial_s": round((first_partial - t0) if first_partial else 0.0, 3),
        "pcm_end_to_final_s": round(final_at - last, 3),
        "wall_s": round(t_end - t0, 3),
    }


CSV_FIELDS = [
    "seq",
    "wav",
    "dur_s",
    "ref",
    "hyp",
    "cer",
    "cer_nopunct",
    "eos_reason",
    "n_partial",
    "first_partial_s",
    "pcm_end_to_final_s",
    "wall_s",
    "session_id",
    "t_start",
    "error",
    "log_kind",
    "log_audio_s",
    "log_queue_wait_ms",
    "log_infer_ms",
    "log_encode_ms",
    "log_decode_ms",
    "log_n_chunks",
    "log_decode_skipped",
    "log_decode_chunks",
]


def _bucket(dur: float) -> str:
    if dur < 8:
        return "<8s"
    if dur < 16:
        return "8-16s"
    if dur < 30:
        return "16-30s"
    return "30-60s"


def _mean(vals: list[float]) -> float:
    return sum(vals) / len(vals) if vals else 0.0


def print_summary(rows: list[dict]) -> None:
    ok = [r for r in rows if r.get("cer_nopunct") != "" and not r.get("error")]
    print()
    print(f"done {len(ok)}/{len(rows)} utterances")
    if not ok:
        return
    print(f"mean CER          { _mean([float(r['cer']) for r in ok]):.4f}")
    print(f"mean CER nopunct  { _mean([float(r['cer_nopunct']) for r in ok]):.4f}")
    print(f"mean pcm→final    { _mean([float(r['pcm_end_to_final_s']) for r in ok]):.3f}s")
    buckets: dict[str, list[dict]] = {}
    for r in ok:
        buckets.setdefault(_bucket(float(r["dur_s"])), []).append(r)
    print("by duration:")
    for name in ("<8s", "8-16s", "16-30s", "30-60s"):
        group = buckets.get(name) or []
        if not group:
            continue
        print(
            f"  {name:7} n={len(group):2}  cer={_mean([float(r['cer_nopunct']) for r in group]):.3f}"
            f"  pcm→final={_mean([float(r['pcm_end_to_final_s']) for r in group]):.3f}s"
            f"  infer_ms={_mean([float(r['log_infer_ms']) for r in group if r.get('log_infer_ms')]) or 0:.0f}"
        )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", type=Path, default=AUDIO_DIR)
    ap.add_argument("--xlsx", type=Path, default=None)
    ap.add_argument("--out", type=Path, default=ROOT / "E2Etest" / "out" / "eval_ws_result.csv")
    ap.add_argument("--log", type=Path, default=None)
    ap.add_argument("--language", default="zh")
    ap.add_argument("--from", dest="seq_from", type=int, default=1)
    ap.add_argument("--to", dest="seq_to", type=int, default=50)
    ap.add_argument("--fast", action="store_true", help="send PCM as fast as possible (accuracy only)")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-vad", action="store_true", help="ignored; server VAD is always on")
    args = ap.parse_args()
    if args.log is None:
        args.log = resolve_server_log()

    folder = args.dir
    xlsx = args.xlsx or (folder / XLSX_NAME)
    if not xlsx.is_file():
        print(f"excel not found: {xlsx}", file=sys.stderr)
        return 1
    try:
        items = read_excel(xlsx)
    except ImportError:
        print("pip install openpyxl", file=sys.stderr)
        return 1
    items = [(s, t) for s, t in items if args.seq_from <= int(s) <= args.seq_to]
    if not items:
        print("no rows in range", file=sys.stderr)
        return 1

    if args.dry_run:
        missing = 0
        for seq, ref in items:
            wav = find_wav(folder, seq)
            if not wav:
                print(f"MISSING {seq}.wav  ref_n={len(ref)}")
                missing += 1
                continue
            _pcm, dur = load_pcm_s16le(wav)
            print(f"{seq:>3}  {dur:6.2f}s  ref_n={len(ref):3}  {wav.name}")
        print(f"dry-run rows={len(items)} missing_wav={missing}")
        return 1 if missing else 0

    try:
        health = httpx.get(f"{BASE_URL}/health", timeout=5.0)
        health.raise_for_status()
    except Exception as exc:
        print(f"server unreachable at {BASE_URL}: {exc}", file=sys.stderr)
        return 1
    print("health:", health.json())
    print(f"sequential {len(items)} files, one WebSocket at a time")

    lock_fd = acquire_eval_lock()
    try:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        results: list[dict] = []
        if args.no_vad:
            print("--no-vad is ignored; server VAD is always on")
        for seq, ref in items:
            wav = find_wav(folder, seq)
            rec = {k: "" for k in CSV_FIELDS}
            rec.update({"seq": seq, "wav": str(wav) if wav else "", "ref": ref})
            if not wav:
                rec["error"] = "MISSING_WAV"
                results.append(rec)
                print(f"{seq} MISSING_WAV")
                continue
            pcm, dur = load_pcm_s16le(wav)
            rec["dur_s"] = round(dur, 3)
            sid = f"eval-{seq}"
            print(f"{seq:>3} start  {dur:5.1f}s  {wav.name}")
            try:
                got = asyncio.run(
                    transcribe_one(
                        pcm,
                        session_id=sid,
                        realtime=not args.fast,
                        chunk_ms=200,
                        language=args.language,
                        timeout_s=180.0,
                    )
                )
                rec.update(got)
                rec["cer"] = round(cer(ref, got["hyp"], False), 4)
                rec["cer_nopunct"] = round(cer(ref, got["hyp"], True), 4)
            except Exception as exc:
                rec["error"] = str(exc)
                rec["session_id"] = sid
            preview = (rec.get("hyp") or rec.get("error") or "")[:40]
            print(
                f"{seq:>3} done   cer={rec.get('cer_nopunct') if rec.get('cer_nopunct') != '' else '-'} "
                f"pcm→final={rec.get('pcm_end_to_final_s') or '-'} "
                f"eos={rec.get('eos_reason') or '-'}  {preview}"
            )
            results.append(rec)
            time.sleep(0.6)

        attach_log_metrics(results, args.log)
        with args.out.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=CSV_FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(results)
        print_summary(results)
        print("wrote", args.out)
        return 0
    finally:
        release_eval_lock(lock_fd)


if __name__ == "__main__":
    raise SystemExit(main())
